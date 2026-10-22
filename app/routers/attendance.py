import json
import math
import re
from io import BytesIO
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import and_, or_
from sqlalchemy.orm import Session

from app.core.security import get_current_user, require_permission
from app.database import get_db
from app.models.attendance import (
    AbsenceRequest,
    AbsenceThreshold,
    AttendanceAudit,
    AttendanceNotification,
    AttendanceRecord,
)
from app.models.chuong_trinh_dao_tao import ChuongTrinhDaoTao
from app.models.diem_danh_buoi_hoc import DiemDanhBuoiHoc
from app.models.lich_hoc import LichHoc
from app.models.lop_hoc import LopHoc
from app.models.lop_hoc_hoc_vien import LopHocHocVien, LopHocHocVienThaoTac
from app.models.nguoi_dung import NguoiDung
from app.models.nguoi_dung_vai_tro import NguoiDungVaiTro
from app.routers.attendance_sessions import (
    create_absence_request,
    decide_absence_request as decide_session_absence_request,
)
from app.schemas.attendance import AbsenceRequestDecision
from app.services.business_time import business_now, business_now_naive
from app.services.attendance_thresholds import (
    absence_metrics_by_student,
    absence_metrics_for_student,
    notify_threshold_activation,
    threshold_for_program as _threshold,
)

router = APIRouter(prefix="/attendance", tags=["Attendance"])
VALID_STATUSES = {"present", "late", "excused_absence", "unexcused_absence"}
MANAGER_ROLES = {7, 8}
TEACHER_ROLES = {3, 4}


class AttendanceItem(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    student_id: int = Field(alias="studentId", gt=0)
    status: str
    reason: str | None = Field(default=None, max_length=1000)


class AttendanceBatch(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    attendance_date: date = Field(alias="date")
    records: list[AttendanceItem]


class AbsenceRequestInput(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    class_id: int = Field(alias="classId", gt=0)
    attendance_date: date = Field(alias="date")
    reason: str = Field(min_length=1, max_length=2000)
    evidence: str | None = Field(default=None, max_length=1000)


class AbsenceDecision(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    decision: str
    reason: str = Field(min_length=1, max_length=1000)


class ThresholdInput(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    percentage: int = Field(alias="percentage", ge=0, le=100)


def _role_ids(db: Session, user: NguoiDung) -> set[int]:
    return {int(row[0]) for row in db.query(NguoiDungVaiTro.VaiTroId).filter(
        NguoiDungVaiTro.NguoiDungId == user.NguoiDungId
    ).all()}


def _manager(db: Session, user: NguoiDung) -> bool:
    return bool(_role_ids(db, user) & MANAGER_ROLES)


def _is_student(db: Session, user: NguoiDung) -> bool:
    return 2 in _role_ids(db, user)


def _class_or_404(db: Session, class_id: int) -> LopHoc:
    record = db.query(LopHoc).filter(LopHoc.LopHocId == class_id).first()
    if record is None:
        raise HTTPException(status_code=404, detail="Class not found")
    return record


def _ensure_teacher_or_manager(db: Session, user: NguoiDung, class_record: LopHoc) -> None:
    roles = _role_ids(db, user)
    if roles & MANAGER_ROLES:
        return
    assigned = class_record.GiangVienId == user.NguoiDungId or db.query(LichHoc.LichHocId).filter(
        LichHoc.LopHocId == class_record.LopHocId,
        LichHoc.GiangVienId == user.NguoiDungId,
    ).first() is not None
    if not (roles & TEACHER_ROLES) or not assigned:
        raise HTTPException(status_code=403, detail="You cannot manage attendance for this class")


def _ensure_student_enrollment(
    db: Session, user: NguoiDung, class_id: int
) -> None:
    enrollment = db.query(LopHocHocVien.LopHocId).filter(
        LopHocHocVien.LopHocId == class_id,
        LopHocHocVien.HocVienId == user.NguoiDungId,
        LopHocHocVien.TrangThaiGhiDanh == "Đang học",
        LopHocHocVien.NgayKetThuc.is_(None),
    ).first()
    if not _is_student(db, user) or enrollment is None:
        raise HTTPException(
            status_code=403,
            detail="Bạn không có tên trong danh sách học viên đang học của lớp này.",
        )


def _ensure_edit_date(attendance_date: date, manager: bool) -> None:
    if attendance_date > date.today():
        raise HTTPException(status_code=400, detail="Attendance can only be recorded for today or earlier")
    if attendance_date < date.today() and not manager:
        raise HTTPException(status_code=403, detail="Only a manager can edit past attendance")


def _audit(db: Session, kind: str, entity_id: int, action: str, old, new, reason, actor_id: int) -> None:
    db.add(AttendanceAudit(
        DoiTuong=kind,
        DoiTuongId=entity_id,
        HanhDong=action,
        DuLieuCu=json.dumps(old, ensure_ascii=False, default=str) if old is not None else None,
        DuLieuMoi=json.dumps(new, ensure_ascii=False, default=str) if new is not None else None,
        LyDo=reason,
        NguoiThucHienId=actor_id,
        ThoiDiem=datetime.now(),
    ))


def _notify_threshold_crossing(db: Session, class_record: LopHoc, student: NguoiDung, before_rate: float) -> None:
    program_id = class_record.ChuongTrinhId
    if program_id is None:
        return
    notify_threshold_activation(
        db,
        student.NguoiDungId,
        program_id,
        was_above_threshold=before_rate >= _threshold(db, program_id),
    )


def _upsert_record(db: Session, class_record: LopHoc, student_id: int, attendance_date: date,
                   state: str, reason: str | None, actor: NguoiDung) -> AttendanceRecord:
    if state not in VALID_STATUSES:
        raise HTTPException(status_code=422, detail=f"Invalid attendance state: {state}")
    if db.query(LopHocHocVien).filter(
        LopHocHocVien.LopHocId == class_record.LopHocId,
        LopHocHocVien.HocVienId == student_id,
        LopHocHocVien.NgayThamGia <= datetime.combine(attendance_date, datetime.max.time()),
        (LopHocHocVien.NgayKetThuc.is_(None))
        | (
            LopHocHocVien.NgayKetThuc
            > datetime.combine(attendance_date, datetime.max.time())
        ),
    ).first() is None:
        raise HTTPException(status_code=400, detail="Student is not enrolled in this class")
    student = db.query(NguoiDung).filter(NguoiDung.NguoiDungId == student_id).first()
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found")
    before_rate = 0.0
    if class_record.ChuongTrinhId is not None:
        before_rate = absence_metrics_for_student(
            db, student_id, class_record.ChuongTrinhId
        )[2]
    record = db.query(AttendanceRecord).filter(
        AttendanceRecord.LopHocId == class_record.LopHocId,
        AttendanceRecord.HocVienId == student_id,
        AttendanceRecord.NgayHoc == attendance_date,
    ).first()
    old = {"status": record.TrangThai, "reason": record.GhiChu} if record else None
    if record is None:
        record = AttendanceRecord(
            LopHocId=class_record.LopHocId,
            HocVienId=student_id,
            NgayHoc=attendance_date,
            TrangThai=state,
            GhiChu=reason,
            NguoiCapNhatId=actor.NguoiDungId,
            NgayCapNhat=datetime.now(),
        )
        db.add(record)
        db.flush()
        action = "created"
    else:
        record.TrangThai = state
        record.GhiChu = reason
        record.NguoiCapNhatId = actor.NguoiDungId
        record.NgayCapNhat = datetime.now()
        action = "updated"
    new = {"status": state, "reason": reason}
    _audit(db, "attendance", record.DiemDanhId, action, old, new, reason, actor.NguoiDungId)
    if state == "unexcused_absence" and old is None or (
        state == "unexcused_absence" and old and old["status"] != state
    ):
        _notify_threshold_crossing(db, class_record, student, before_rate)
    return record


def _class_payload(db: Session, class_record: LopHoc) -> dict:
    return {
        "classId": class_record.LopHocId,
        "classCode": class_record.MaLop,
        "className": class_record.TenLop,
        "programId": class_record.ChuongTrinhId,
        "absenceThreshold": _threshold(db, class_record.ChuongTrinhId),
    }


@router.get("/classes")
def list_attendance_classes(
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db)
):
    roles = _role_ids(db, current_user)
    query = db.query(LopHoc)
    if _is_student(db, current_user):
        enrolled = db.query(LopHocHocVien.LopHocId).filter(
            LopHocHocVien.HocVienId == current_user.NguoiDungId,
            LopHocHocVien.TrangThaiGhiDanh == "Đang học",
            LopHocHocVien.NgayKetThuc.is_(None),
        )
        query = query.filter(LopHoc.LopHocId.in_(enrolled))
    elif not roles & MANAGER_ROLES:
        if roles & TEACHER_ROLES:
            taught = db.query(LichHoc.LopHocId).filter(LichHoc.GiangVienId == current_user.NguoiDungId)
            query = query.filter(or_(LopHoc.GiangVienId == current_user.NguoiDungId, LopHoc.LopHocId.in_(taught)))
        else:
            raise HTTPException(status_code=403, detail="Attendance access is not allowed")
    return [_class_payload(db, record) for record in query.order_by(LopHoc.MaLop).all()]


@router.get("/classes/{class_id}/matrix")
def get_attendance_matrix(
    class_id: int,
    attendance_date: date | None = Query(default=None, alias="date"),
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    _ensure_teacher_or_manager(db, current_user, class_record)
    day = attendance_date or date.today()
    roster = db.query(NguoiDung).join(
        LopHocHocVien, LopHocHocVien.HocVienId == NguoiDung.NguoiDungId
    ).filter(
        LopHocHocVien.LopHocId == class_id,
        LopHocHocVien.NgayThamGia <= datetime.combine(day, datetime.max.time()),
        (LopHocHocVien.NgayKetThuc.is_(None))
        | (LopHocHocVien.NgayKetThuc > datetime.combine(day, datetime.max.time())),
    ).order_by(NguoiDung.HoTen).all()
    recorded = {row.HocVienId: row for row in db.query(AttendanceRecord).filter(
        AttendanceRecord.LopHocId == class_id, AttendanceRecord.NgayHoc == day
    ).all()}
    requests = db.query(AbsenceRequest).filter(
        AbsenceRequest.LopHocId == class_id, AbsenceRequest.NgayHoc == day
    ).all()
    requests_by_student = {request.HocVienId: request for request in requests if request.TrangThai == "pending"}
    rows = []
    for student in roster:
        record = recorded.get(student.NguoiDungId)
        request = requests_by_student.get(student.NguoiDungId)
        rows.append({
            "studentId": student.NguoiDungId,
            "studentName": student.HoTen,
            "email": student.Email,
            "status": record.TrangThai if record else "unmarked",
            "reason": record.GhiChu if record else None,
            "requestId": request.DonXinVangId if request else None,
        })
    present = sum(1 for row in rows if row["status"] == "present")
    late = sum(1 for row in rows if row["status"] == "late")
    excused = sum(1 for row in rows if row["status"] == "excused_absence")
    unmarked = sum(1 for row in rows if row["status"] == "unmarked")
    absent = sum(1 for row in rows if row["status"] in {"excused_absence", "unexcused_absence"})
    unexcused = sum(1 for row in rows if row["status"] == "unexcused_absence")
    threshold = _threshold(db, class_record.ChuongTrinhId)
    flagged_students = []
    if class_record.ChuongTrinhId is not None:
        metrics = absence_metrics_by_student(
            db,
            class_record.ChuongTrinhId,
            (student.NguoiDungId for student in roster),
        )
        for student in roster:
            total_absent, total_marked, rate = metrics.get(
                student.NguoiDungId, (0, 0, 0.0)
            )
            if total_marked and rate >= threshold:
                flagged_students.append({"studentId": student.NguoiDungId, "studentName": student.HoTen,
                                         "absencePercent": round(rate, 1), "unexcused": total_absent,
                                         "total": total_marked})
    flagged_by_student = {
        item["studentId"]: item for item in flagged_students
    }
    for row in rows:
        absent, total, rate = (
            metrics.get(row["studentId"], (0, 0, 0.0))
            if class_record.ChuongTrinhId is not None
            else (0, 0, 0.0)
        )
        row["absencePercent"] = round(rate, 1)
        row["unexcusedAbsences"] = absent
        row["totalMarked"] = total
        row["absenceThreshold"] = threshold
        row["aboveThreshold"] = row["studentId"] in flagged_by_student
    return {
        **_class_payload(db, class_record),
        "date": day.isoformat(),
        "students": rows,
        "counts": {
            "present": present,
            "late": late,
            "excused": excused,
            "unexcused": unexcused,
            "absent": absent,
            "unmarked": unmarked,
            "total": len(rows),
        },
        "attendancePercent": round((present + late) * 100 / len(rows), 1) if rows else 0,
        "absencePercent": round(unexcused * 100 / len(rows), 1) if rows else 0,
        "flaggedStudents": flagged_students,
        "aboveThreshold": bool(flagged_students),
    }


@router.post("/classes/{class_id}/records")
def save_attendance(
    class_id: int,
    payload: AttendanceBatch,
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    _ensure_teacher_or_manager(db, current_user, class_record)
    manager = _manager(db, current_user)
    _ensure_edit_date(payload.attendance_date, manager)
    touched = set()
    try:
        for item in payload.records:
            if item.student_id in touched:
                raise HTTPException(status_code=422, detail="A student can only appear once per attendance batch")
            touched.add(item.student_id)
            _upsert_record(db, class_record, item.student_id, payload.attendance_date,
                           item.status, item.reason, current_user)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return get_attendance_matrix(class_id, payload.attendance_date, current_user, db)


@router.post("/classes/{class_id}/mark-all-present")
def mark_all_present(
    class_id: int,
    attendance_date: date = Query(default_factory=date.today, alias="date"),
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    _ensure_teacher_or_manager(db, current_user, class_record)
    _ensure_edit_date(attendance_date, _manager(db, current_user))
    students = db.query(LopHocHocVien.HocVienId).filter(
        LopHocHocVien.LopHocId == class_id,
        LopHocHocVien.NgayThamGia <= datetime.combine(attendance_date, datetime.max.time()),
        (LopHocHocVien.NgayKetThuc.is_(None))
        | (LopHocHocVien.NgayKetThuc > datetime.combine(attendance_date, datetime.max.time())),
    ).all()
    try:
        for (student_id,) in students:
            _upsert_record(db, class_record, student_id, attendance_date, "present", None, current_user)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return get_attendance_matrix(class_id, attendance_date, current_user, db)


@router.get("/classes/{class_id}/audit")
def list_attendance_audit(
    class_id: int,
    attendance_date: date | None = Query(default=None, alias="date"),
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    _ensure_teacher_or_manager(db, current_user, class_record)
    ids = [row[0] for row in db.query(AttendanceRecord.DiemDanhId).filter(
        AttendanceRecord.LopHocId == class_id,
        *([AttendanceRecord.NgayHoc == attendance_date] if attendance_date else []),
    ).all()]
    audit = db.query(AttendanceAudit).filter(
        AttendanceAudit.DoiTuong == "attendance",
        AttendanceAudit.DoiTuongId.in_(ids),
    ).order_by(AttendanceAudit.ThoiDiem.desc()).all() if ids else []
    return [{
        "id": row.LichSuId, "action": row.HanhDong, "old": row.DuLieuCu,
        "new": row.DuLieuMoi, "reason": row.LyDo, "at": row.ThoiDiem.isoformat(),
        "actorId": row.NguoiThucHienId,
    } for row in audit]


@router.post("/absence-requests")
def submit_absence_request(
    payload: AbsenceRequestInput,
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    enrolled = db.query(LopHocHocVien).filter(
        LopHocHocVien.LopHocId == payload.class_id,
        LopHocHocVien.HocVienId == current_user.NguoiDungId,
        LopHocHocVien.NgayThamGia
        <= datetime.combine(payload.attendance_date, datetime.max.time()),
        (LopHocHocVien.NgayKetThuc.is_(None))
        | (
            LopHocHocVien.NgayKetThuc
            > datetime.combine(payload.attendance_date, datetime.max.time())
        ),
    ).first()
    if enrolled is None or not _is_student(db, current_user):
        raise HTTPException(status_code=403, detail="Only an enrolled student can request an excused absence")
    if payload.attendance_date > date.today():
        scheduled = db.query(LichHoc).filter(
            LichHoc.LopHocId == payload.class_id,
            LichHoc.NgayHoc == payload.attendance_date,
            LichHoc.TrangThai != "Đã hủy",
        ).first()
        if scheduled is None:
            raise HTTPException(status_code=400, detail="No scheduled class exists on this date")
    pending = db.query(AbsenceRequest).filter(
        AbsenceRequest.LopHocId == payload.class_id,
        AbsenceRequest.HocVienId == current_user.NguoiDungId,
        AbsenceRequest.NgayHoc == payload.attendance_date,
        AbsenceRequest.TrangThai == "pending",
    ).first()
    if pending:
        raise HTTPException(status_code=409, detail="A pending request already exists for this class and date")
    request = AbsenceRequest(
        LopHocId=payload.class_id,
        HocVienId=current_user.NguoiDungId,
        NgayHoc=payload.attendance_date,
        LyDo=payload.reason.strip(),
        TepChungMinh=payload.evidence,
        TrangThai="pending",
        NgayTao=datetime.now(),
    )
    if not request.LyDo:
        raise HTTPException(status_code=422, detail="A reason is required")
    db.add(request)
    db.flush()
    _audit(db, "absence_request", request.DonXinVangId, "submitted", None,
           {"status": "pending", "reason": request.LyDo, "evidence": request.TepChungMinh},
           None, current_user.NguoiDungId)
    db.commit()
    return {"requestId": request.DonXinVangId, "status": request.TrangThai, "date": request.NgayHoc.isoformat()}


@router.get("/absence-requests")
def list_absence_requests(
    status: str | None = Query(default=None),
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    roles = _role_ids(db, current_user)
    query = db.query(AbsenceRequest)
    if roles & MANAGER_ROLES:
        pass
    elif roles & TEACHER_ROLES:
        taught = db.query(LopHoc.LopHocId).filter(LopHoc.GiangVienId == current_user.NguoiDungId)
        scheduled = db.query(LichHoc.LopHocId).filter(LichHoc.GiangVienId == current_user.NguoiDungId)
        query = query.filter(or_(AbsenceRequest.LopHocId.in_(taught), AbsenceRequest.LopHocId.in_(scheduled)))
    elif 2 in roles:
        query = query.filter(AbsenceRequest.HocVienId == current_user.NguoiDungId)
    else:
        raise HTTPException(status_code=403, detail="Attendance access is not allowed")
    if status:
        query = query.filter(AbsenceRequest.TrangThai == status)
    result = []
    for request in query.order_by(AbsenceRequest.NgayTao.desc()).all():
        class_record = db.query(LopHoc).filter(LopHoc.LopHocId == request.LopHocId).first()
        student = db.query(NguoiDung).filter(NguoiDung.NguoiDungId == request.HocVienId).first()
        result.append({
            "requestId": request.DonXinVangId, "classId": request.LopHocId,
            "classCode": class_record.MaLop if class_record else "",
            "studentId": request.HocVienId, "studentName": student.HoTen if student else "",
            "date": request.NgayHoc.isoformat(), "reason": request.LyDo,
            "evidence": request.TepChungMinh, "status": request.TrangThai,
            "reviewReason": request.LyDoXuLy,
        })
    return result


@router.post("/absence-requests/{request_id}/decision")
def decide_class_absence_request(
    request_id: int,
    payload: AbsenceDecision,
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    if payload.decision not in {"approved", "rejected"}:
        raise HTTPException(status_code=422, detail="Decision must be approved or rejected")
    if not payload.reason.strip():
        raise HTTPException(status_code=422, detail="A review reason is required")
    request = db.query(AbsenceRequest).filter(AbsenceRequest.DonXinVangId == request_id).first()
    if request is None:
        raise HTTPException(status_code=404, detail="Absence request not found")
    class_record = _class_or_404(db, request.LopHocId)
    _ensure_teacher_or_manager(db, current_user, class_record)
    if request.TrangThai != "pending":
        raise HTTPException(status_code=409, detail="This absence request has already been reviewed")
    old = {"status": request.TrangThai, "reviewReason": request.LyDoXuLy}
    request.TrangThai = payload.decision
    request.LyDoXuLy = payload.reason.strip()
    request.NguoiXuLyId = current_user.NguoiDungId
    request.NgayXuLy = datetime.now()
    _audit(db, "absence_request", request.DonXinVangId, payload.decision, old,
           {"status": request.TrangThai, "reviewReason": request.LyDoXuLy}, request.LyDoXuLy, current_user.NguoiDungId)
    if payload.decision == "approved":
        _upsert_record(db, class_record, request.HocVienId, request.NgayHoc,
                       "excused_absence", request.LyDo, current_user)
    db.add(AttendanceNotification(
        NguoiNhanId=request.HocVienId, HocVienId=request.HocVienId,
        ChuongTrinhId=class_record.ChuongTrinhId, Loai="absence_request_" + payload.decision,
        TieuDe="Absence request " + payload.decision,
        NoiDung=f"Your request for {request.NgayHoc.isoformat()} was {payload.decision}. {request.LyDoXuLy}",
        ThoiDiem=datetime.now(),
    ))
    db.commit()
    return {"requestId": request.DonXinVangId, "status": request.TrangThai, "reviewReason": request.LyDoXuLy}


@router.get("/my-history")
def get_my_attendance_history(
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    if not _is_student(db, current_user):
        raise HTTPException(status_code=403, detail="This endpoint is for students")
    daily_rows = db.query(AttendanceRecord, LopHoc).join(
        LopHoc, LopHoc.LopHocId == AttendanceRecord.LopHocId
    ).filter(AttendanceRecord.HocVienId == current_user.NguoiDungId).order_by(
        AttendanceRecord.NgayHoc.desc()
    ).all()
    daily_by_class_date = {
        (class_record.LopHocId, record.NgayHoc): record
        for record, class_record in daily_rows
    }
    daily_status_labels = {
        "present": "Có mặt",
        "late": "Đi muộn",
        "excused_absence": "Vắng có phép",
        "unexcused_absence": "Vắng không phép",
    }
    session_status_labels = {
        "Có mặt": "Có mặt",
        "Đi muộn": "Đi muộn",
        "Vắng có phép": "Vắng có phép",
        "Vắng không phép": "Vắng không phép",
    }
    now = business_now_naive()
    schedule_rows = (
        db.query(LichHoc, LopHoc, LopHocHocVien)
        .join(LopHoc, LopHoc.LopHocId == LichHoc.LopHocId)
        .join(LopHocHocVien, LopHocHocVien.LopHocId == LopHoc.LopHocId)
        .filter(
            LopHocHocVien.HocVienId == current_user.NguoiDungId,
            LopHocHocVien.NgayThamGia
            < datetime.combine(now.date() + timedelta(days=1), datetime.min.time()),
            LichHoc.TrangThai == "Đã lên lịch",
            or_(
                LichHoc.NgayHoc < now.date(),
                and_(
                    LichHoc.NgayHoc == now.date(),
                    LichHoc.GioBatDau <= now.time(),
                ),
            ),
        )
        .order_by(LichHoc.NgayHoc.desc(), LichHoc.GioBatDau.desc())
        .all()
    )
    ended_dates: dict[int, list[date]] = {}
    for class_id, effective_date in db.query(
        LopHocHocVienThaoTac.LopHocId,
        LopHocHocVienThaoTac.NgayHieuLuc,
    ).filter(
        LopHocHocVienThaoTac.HocVienId == current_user.NguoiDungId
    ).all():
        ended_dates.setdefault(class_id, []).append(effective_date)
    eligible_schedules = [
        (schedule, class_record)
        for schedule, class_record, enrollment in schedule_rows
        if enrollment.NgayThamGia
        < datetime.combine(
            schedule.NgayHoc + timedelta(days=1),
            datetime.min.time(),
        )
        and (
            enrollment.NgayKetThuc is None
            or enrollment.NgayKetThuc
            > datetime.combine(schedule.NgayHoc, datetime.min.time())
        )
        and not any(
            effective_date <= schedule.NgayHoc
            for effective_date in ended_dates.get(class_record.LopHocId, [])
        )
    ]
    schedule_ids = [schedule.LichHocId for schedule, _ in eligible_schedules]
    session_attendance = {
        (item.LichHocId, item.HocVienId): item
        for item in db.query(DiemDanhBuoiHoc).filter(
            DiemDanhBuoiHoc.HocVienId == current_user.NguoiDungId,
            DiemDanhBuoiHoc.LichHocId.in_(schedule_ids),
        ).all()
    } if schedule_ids else {}
    recorded_sessions_by_class_date: set[tuple[int, date]] = set()
    for schedule, class_record in eligible_schedules:
        key = (class_record.LopHocId, schedule.NgayHoc)
        if (schedule.LichHocId, current_user.NguoiDungId) in session_attendance:
            recorded_sessions_by_class_date.add(key)

    history = []
    summaries: dict[int, dict] = {}
    for record, class_record in daily_rows:
        key = (class_record.LopHocId, record.NgayHoc)
        if key in recorded_sessions_by_class_date:
            continue
        history.append({
            "classId": class_record.LopHocId,
            "classCode": class_record.MaLop,
            "className": class_record.TenLop,
            "programId": class_record.ChuongTrinhId,
            "date": record.NgayHoc.isoformat(),
            "startTime": None,
            "endTime": None,
            "scheduleCode": None,
            "status": daily_status_labels.get(record.TrangThai, record.TrangThai),
            "reason": record.GhiChu,
        })
    for schedule, class_record in eligible_schedules:
        key = (class_record.LopHocId, schedule.NgayHoc)
        if key in daily_by_class_date and key not in recorded_sessions_by_class_date:
            continue
        attendance = session_attendance.get(
            (schedule.LichHocId, current_user.NguoiDungId)
        )
        history.append({
            "classId": class_record.LopHocId,
            "classCode": class_record.MaLop,
            "className": class_record.TenLop,
            "programId": class_record.ChuongTrinhId,
            "date": schedule.NgayHoc.isoformat(),
            "startTime": schedule.GioBatDau.strftime("%H:%M"),
            "endTime": schedule.GioKetThuc.strftime("%H:%M"),
            "scheduleCode": f"BH-{schedule.LichHocId:06d}",
            "status": session_status_labels.get(
                attendance.TrangThai, attendance.TrangThai
            )
            if attendance
            else "Chưa điểm danh",
            "reason": attendance.GhiChu if attendance else None,
        })
    history.sort(
        key=lambda item: (
            item["date"],
            item["startTime"] or "",
            item["classId"],
        ),
        reverse=True,
    )
    for item in history:
        summary = summaries.setdefault(
            item["classId"],
            {
                "classId": item["classId"],
                "classCode": item["classCode"],
                "className": item["className"],
                "programId": item["programId"],
                "present": 0,
                "late": 0,
                "excusedAbsences": 0,
                "unexcusedAbsences": 0,
                "unmarked": 0,
                "totalMarked": 0,
            },
        )
        status_value = item["status"]
        status_counts = {
            "Có mặt": "present",
            "Đi muộn": "late",
            "Vắng có phép": "excusedAbsences",
            "Vắng không phép": "unexcusedAbsences",
            "Chưa điểm danh": "unmarked",
        }
        count_key = status_counts.get(status_value)
        if count_key:
            summary[count_key] += 1
        if status_value != "Chưa điểm danh":
            summary["totalMarked"] += 1
    attendance_summaries = []
    for summary in summaries.values():
        total = summary["totalMarked"]
        absent = summary["unexcusedAbsences"]
        threshold = _threshold(db, summary["programId"])
        absence_percent = absent * 100 / total if total else 0
        threshold_reached = total > 0 and absence_percent >= threshold
        if total == 0:
            remaining_absences = None
        elif threshold_reached or threshold <= 0:
            remaining_absences = 0
        elif threshold >= 100 and absent < total:
            remaining_absences = None
        else:
            absences_until_threshold = math.ceil(
                (threshold * total / 100 - absent) / (1 - threshold / 100)
            )
            remaining_absences = max(0, absences_until_threshold - 1)
        summary.update({
            "totalSessions": total + summary["unmarked"],
            "absencePercent": round(absence_percent, 1),
            "attendancePercent": round(
                (summary["present"] + summary["late"])
                * 100
                / (total + summary["unmarked"]),
                1,
            )
            if total + summary["unmarked"]
            else 0,
            "threshold": threshold,
            "thresholdReached": threshold_reached,
            "remainingAbsences": remaining_absences,
            "nearThreshold": (
                not threshold_reached
                and remaining_absences is not None
                and remaining_absences <= 1
            ),
        })
        attendance_summaries.append(summary)
    attendance_summaries.sort(key=lambda item: item["classCode"])
    warning_summary = next(
        (item for item in attendance_summaries if item["thresholdReached"]),
        None,
    )
    warning = (
        {
            "programId": warning_summary["programId"],
            "classId": warning_summary["classId"],
            "absencePercent": warning_summary["absencePercent"],
            "threshold": warning_summary["threshold"],
            "message": (
                f"Tỷ lệ vắng không phép của bạn ở lớp "
                f"{warning_summary['classCode']} là "
                f"{warning_summary['absencePercent']:.1f}% "
                f"({warning_summary['unexcusedAbsences']}/"
                f"{warning_summary['totalMarked']}), đạt ngưỡng "
                f"{warning_summary['threshold']}%."
            ),
        }
        if warning_summary
        else None
    )
    requests = list_absence_requests(None, current_user, db)
    notifications = db.query(AttendanceNotification).filter(
        AttendanceNotification.NguoiNhanId == current_user.NguoiDungId
    ).order_by(AttendanceNotification.ThoiDiem.desc()).limit(50).all()
    notices = [{"id": item.ThongBaoId, "type": item.Loai, "title": item.TieuDe,
                "message": item.NoiDung, "at": item.ThoiDiem.isoformat(), "read": item.DaDoc is not None}
               for item in notifications]
    return {
        "records": history,
        "requests": requests,
        "warning": warning,
        "attendanceSummaries": attendance_summaries,
        "notifications": notices,
    }


@router.get("/notifications")
def list_attendance_notifications(
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    items = db.query(AttendanceNotification).filter(
        AttendanceNotification.NguoiNhanId == current_user.NguoiDungId
    ).order_by(AttendanceNotification.ThoiDiem.desc()).limit(50).all()
    return [{"id": item.ThongBaoId, "type": item.Loai, "title": item.TieuDe,
             "message": item.NoiDung, "at": item.ThoiDiem.isoformat(), "read": item.DaDoc is not None}
            for item in items]


@router.get("/thresholds")
def list_thresholds(
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    if not _manager(db, current_user):
        raise HTTPException(status_code=403, detail="Only a manager can view program thresholds")
    return [{"programId": row.ChuongTrinhId, "programCode": row.MaChuongTrinh,
             "programName": row.TenChuongTrinh, "percentage": _threshold(db, row.ChuongTrinhId)}
            for row in db.query(ChuongTrinhDaoTao).order_by(ChuongTrinhDaoTao.MaChuongTrinh).all()]


@router.put("/thresholds/{program_id}")
def update_threshold(
    program_id: int,
    payload: ThresholdInput,
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    if not _manager(db, current_user):
        raise HTTPException(status_code=403, detail="Only a manager can update program thresholds")
    if db.query(ChuongTrinhDaoTao).filter(ChuongTrinhDaoTao.ChuongTrinhId == program_id).first() is None:
        raise HTTPException(status_code=404, detail="Program not found")
    row = db.query(AbsenceThreshold).filter(AbsenceThreshold.ChuongTrinhId == program_id).first()
    old_threshold = row.TiLeToiDa if row else 20
    old = {"percentage": old_threshold}
    if row is None:
        row = AbsenceThreshold(ChuongTrinhId=program_id)
        db.add(row)
    row.TiLeToiDa = payload.percentage
    row.NguoiCapNhatId = current_user.NguoiDungId
    row.NgayCapNhat = datetime.now()
    _audit(db, "threshold", program_id, "updated", old, {"percentage": payload.percentage}, None,
           current_user.NguoiDungId)
    db.flush()
    for student_id, (_absent, _total, rate) in absence_metrics_by_student(
        db, program_id
    ).items():
        notify_threshold_activation(
            db,
            student_id,
            program_id,
            was_above_threshold=rate >= old_threshold,
        )
    db.commit()
    return {"programId": program_id, "percentage": row.TiLeToiDa}


@router.get("/classes/{class_id}/export")
def export_attendance(
    class_id: int,
    attendance_date: date | None = Query(default=None, alias="date"),
    current_user: NguoiDung = Depends(get_current_user), db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    _ensure_teacher_or_manager(db, current_user, class_record)
    day = attendance_date or date.today()
    matrix = get_attendance_matrix(class_id, day, current_user, db)
    labels = {
        "present": "Có mặt",
        "late": "Đi muộn",
        "excused_absence": "Vắng có phép",
        "unexcused_absence": "Vắng không phép",
        "unmarked": "Chưa điểm danh",
    }
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tổng hợp"
    counts = matrix["counts"]
    summary.append(["Lớp", class_record.MaLop, class_record.TenLop])
    summary.append(["Ngày", day.isoformat()])
    summary.append(["Tổng học viên", counts["total"]])
    summary.append(["Có mặt", counts["present"]])
    summary.append(["Đi muộn", counts["late"]])
    summary.append(["Vắng có phép", counts["excused"]])
    summary.append(["Vắng không phép", counts["unexcused"]])
    summary.append(["Chưa điểm danh", counts["unmarked"]])
    summary.append(["Tỷ lệ tham gia (%)", matrix["attendancePercent"]])
    details = workbook.create_sheet("Chi tiết")
    details.append(["Mã học viên", "Họ tên", "Email", "Ngày", "Trạng thái", "Ghi chú"])
    for row in matrix["students"]:
        details.append([
            row["studentId"],
            row["studentName"],
            row["email"],
            day.isoformat(),
            labels[row["status"]],
            row["reason"] or "",
        ])
    output = BytesIO()
    workbook.save(output)
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", class_record.MaLop)
    filename = f"{safe_code}-{day.isoformat()}.xlsx"
    return Response(content=output.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get(
    "/classes/{class_id}/grade-matrix",
    dependencies=[Depends(require_permission("ATTENDANCE_VIEW"))],
)
def get_attendance_grade_matrix(
    class_id: int,
    current_user: NguoiDung = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    student_mode = _is_student(db, current_user)
    if student_mode:
        _ensure_student_enrollment(db, current_user, class_id)
    else:
        _ensure_teacher_or_manager(db, current_user, class_record)
    now = business_now_naive()
    schedules = (
        db.query(LichHoc)
        .filter(
            LichHoc.LopHocId == class_id,
            LichHoc.TrangThai == "Đã lên lịch",
            or_(
                LichHoc.NgayHoc < now.date(),
                and_(
                    LichHoc.NgayHoc == now.date(),
                    LichHoc.GioBatDau <= now.time(),
                ),
            ),
        )
        .order_by(LichHoc.NgayHoc, LichHoc.GioBatDau, LichHoc.LichHocId)
        .all()
    )
    enrollment_query = (
        db.query(LopHocHocVien, NguoiDung)
        .join(NguoiDung, NguoiDung.NguoiDungId == LopHocHocVien.HocVienId)
        .filter(LopHocHocVien.LopHocId == class_id)
    )
    if student_mode:
        enrollment_query = enrollment_query.filter(
            LopHocHocVien.HocVienId == current_user.NguoiDungId,
            LopHocHocVien.TrangThaiGhiDanh == "Đang học",
            LopHocHocVien.NgayKetThuc.is_(None),
        )
    enrollments = enrollment_query.order_by(
        NguoiDung.HoTen, NguoiDung.NguoiDungId
    ).all()
    enrollment_by_student = {
        enrollment.HocVienId: enrollment for enrollment, _ in enrollments
    }
    ended_dates: dict[int, list[date]] = {}
    for student_id, effective_date in db.query(
        LopHocHocVienThaoTac.HocVienId,
        LopHocHocVienThaoTac.NgayHieuLuc,
    ).filter(LopHocHocVienThaoTac.LopHocId == class_id).all():
        ended_dates.setdefault(student_id, []).append(effective_date)
    attendance = {
        (schedule_id, student_id): status_value
        for schedule_id, student_id, status_value in (
            db.query(
                DiemDanhBuoiHoc.LichHocId,
                DiemDanhBuoiHoc.HocVienId,
                DiemDanhBuoiHoc.TrangThai,
            )
            .join(LichHoc, LichHoc.LichHocId == DiemDanhBuoiHoc.LichHocId)
            .filter(
                LichHoc.LopHocId == class_id,
                LichHoc.TrangThai == "Đã lên lịch",
                LichHoc.LichHocId.in_(
                    [schedule.LichHocId for schedule in schedules]
                ),
            )
            .all()
            if schedules
            else []
        )
    }
    status_labels = {
        "Có mặt": "Có mặt",
        "Đi muộn": "Đi muộn",
        "Vắng có phép": "Vắng có phép",
        "Vắng không phép": "Vắng không phép",
    }
    status_values = (
        "Có mặt",
        "Đi muộn",
        "Vắng có phép",
        "Vắng không phép",
        "Chưa điểm danh",
    )
    totals = {label: 0 for label in status_values}
    students = []
    for enrollment, student in enrollments:
        counts = {label: 0 for label in status_values}
        session_statuses = []
        total_sessions = 0
        for schedule in schedules:
            enrolled_by_date = enrollment.NgayThamGia < datetime.combine(
                schedule.NgayHoc + timedelta(days=1),
                datetime.min.time(),
            )
            has_ended = any(
                effective_date <= schedule.NgayHoc
                for effective_date in ended_dates.get(student.NguoiDungId, [])
            )
            if not enrolled_by_date or has_ended:
                session_statuses.append("—")
                continue
            status_label = status_labels.get(
                attendance.get((schedule.LichHocId, student.NguoiDungId)),
                "Chưa điểm danh",
            )
            session_statuses.append(status_label)
            counts[status_label] += 1
            totals[status_label] += 1
            total_sessions += 1
        students.append(
            {
                "studentId": student.NguoiDungId,
                "studentName": student.HoTen,
                "email": student.Email or "",
                "attendance": session_statuses,
                "totalSessions": total_sessions,
                "totalAbsences": counts["Vắng có phép"] + counts["Vắng không phép"],
                "present": counts["Có mặt"],
                "late": counts["Đi muộn"],
                "excusedAbsences": counts["Vắng có phép"],
                "unexcusedAbsences": counts["Vắng không phép"],
                "unmarked": counts["Chưa điểm danh"],
                "attendanceRate": round(
                    (counts["Có mặt"] + counts["Đi muộn"]) * 100 / total_sessions,
                    1,
                )
                if total_sessions
                else 0,
            }
        )
    students.sort(
        key=lambda item: (item["attendanceRate"], item["studentName"].casefold())
    )
    return {
        "classId": class_record.LopHocId,
        "classCode": class_record.MaLop,
        "className": class_record.TenLop,
        "sessions": [
            {
                "scheduleId": schedule.LichHocId,
                "scheduleCode": f"BH-{schedule.LichHocId:06d}",
                "date": schedule.NgayHoc.isoformat(),
                "startTime": schedule.GioBatDau.strftime("%H:%M"),
                "endTime": schedule.GioKetThuc.strftime("%H:%M"),
            }
            for schedule in schedules
        ],
        "students": students,
        "totals": totals,
    }


@router.get(
    "/classes/{class_id}/export-matrix",
    dependencies=[Depends(require_permission("ATTENDANCE_VIEW"))],
)
def export_attendance_matrix(
    class_id: int,
    current_user: NguoiDung = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    class_record = _class_or_404(db, class_id)
    student_mode = _is_student(db, current_user)
    if student_mode:
        _ensure_student_enrollment(db, current_user, class_id)
    else:
        _ensure_teacher_or_manager(db, current_user, class_record)
    now = business_now_naive()
    schedules = (
        db.query(LichHoc)
        .filter(
            LichHoc.LopHocId == class_id,
            LichHoc.TrangThai == "Đã lên lịch",
            or_(
                LichHoc.NgayHoc < now.date(),
                and_(
                    LichHoc.NgayHoc == now.date(),
                    LichHoc.GioBatDau <= now.time(),
                ),
            ),
        )
        .order_by(LichHoc.NgayHoc, LichHoc.GioBatDau, LichHoc.LichHocId)
        .all()
    )
    enrollment_query = (
        db.query(LopHocHocVien, NguoiDung)
        .join(NguoiDung, NguoiDung.NguoiDungId == LopHocHocVien.HocVienId)
        .filter(LopHocHocVien.LopHocId == class_id)
    )
    if student_mode:
        enrollment_query = enrollment_query.filter(
            LopHocHocVien.HocVienId == current_user.NguoiDungId
        )
    enrollments = enrollment_query.order_by(
        NguoiDung.HoTen, NguoiDung.NguoiDungId
    ).all()
    students = {student.NguoiDungId: student for _, student in enrollments}
    enrollment_by_student = {
        enrollment.HocVienId: enrollment for enrollment, _ in enrollments
    }
    actions = (
        db.query(
            LopHocHocVienThaoTac.HocVienId,
            LopHocHocVienThaoTac.NgayHieuLuc,
        )
        .filter(LopHocHocVienThaoTac.LopHocId == class_id)
        .all()
    )
    ended_dates: dict[int, list[date]] = {}
    for student_id, effective_date in actions:
        ended_dates.setdefault(student_id, []).append(effective_date)
    attendance_rows = (
        db.query(
            DiemDanhBuoiHoc.LichHocId,
            DiemDanhBuoiHoc.HocVienId,
            DiemDanhBuoiHoc.TrangThai,
        )
        .join(LichHoc, LichHoc.LichHocId == DiemDanhBuoiHoc.LichHocId)
        .filter(
            LichHoc.LopHocId == class_id,
            LichHoc.TrangThai == "Đã lên lịch",
            LichHoc.LichHocId.in_([schedule.LichHocId for schedule in schedules]),
        )
        .all()
        if schedules
        else []
    )
    attendance = {
        (schedule_id, student_id): status_value
        for schedule_id, student_id, status_value in attendance_rows
    }
    status_labels = {
        "Có mặt": "Có mặt",
        "Đi muộn": "Đi muộn",
        "Vắng có phép": "Vắng có phép",
        "Vắng không phép": "Vắng không phép",
    }
    status_fills = {
        "Có mặt": PatternFill("solid", fgColor="DCFCE7"),
        "Đi muộn": PatternFill("solid", fgColor="FEF3C7"),
        "Vắng có phép": PatternFill("solid", fgColor="DBEAFE"),
        "Vắng không phép": PatternFill("solid", fgColor="FEE2E2"),
        "Chưa điểm danh": PatternFill("solid", fgColor="E2E8F0"),
        "—": PatternFill("solid", fgColor="F8FAFC"),
    }
    workbook = Workbook()
    matrix_sheet = workbook.active
    matrix_sheet.title = "Ma trận chuyên cần"
    headers = ["Mã học viên", "Họ tên", "Email"]
    headers.extend(
        f"BH-{schedule.LichHocId:06d}\n"
        f"{schedule.NgayHoc.isoformat()} "
        f"{schedule.GioBatDau.strftime('%H:%M')}-"
        f"{schedule.GioKetThuc.strftime('%H:%M')}"
        for schedule in schedules
    )
    headers.extend(
        [
            "Tổng buổi",
            "Tổng buổi vắng",
            "Có mặt",
            "Đi muộn",
            "Vắng có phép",
            "Vắng không phép",
            "Chưa điểm danh",
            "Tỷ lệ chuyên cần (%)",
        ]
    )
    matrix_sheet.append(headers)
    totals = {
        "Có mặt": 0,
        "Đi muộn": 0,
        "Vắng có phép": 0,
        "Vắng không phép": 0,
        "Chưa điểm danh": 0,
    }
    student_rows = []
    for student_id, student in students.items():
        row = [student_id, student.HoTen, student.Email or ""]
        counts = {label: 0 for label in totals}
        total_sessions = 0
        for schedule in schedules:
            enrollment = enrollment_by_student[student_id]
            enrolled_by_date = (
                enrollment.NgayThamGia
                < datetime.combine(
                    schedule.NgayHoc + timedelta(days=1),
                    datetime.min.time(),
                )
            )
            has_ended = any(
                effective_date <= schedule.NgayHoc
                for effective_date in ended_dates.get(student_id, [])
            )
            if not enrolled_by_date or has_ended:
                row.append("—")
                continue
            total_sessions += 1
            status_label = status_labels.get(
                attendance.get((schedule.LichHocId, student_id)),
                "Chưa điểm danh",
            )
            row.append(status_label)
            counts[status_label] += 1
            totals[status_label] += 1
        attendance_rate = (
            round(
                (counts["Có mặt"] + counts["Đi muộn"])
                * 100
                / total_sessions,
                1,
            )
            if total_sessions
            else 0
        )
        row.extend(
            [
                total_sessions,
                counts["Vắng có phép"] + counts["Vắng không phép"],
                counts["Có mặt"],
                counts["Đi muộn"],
                counts["Vắng có phép"],
                counts["Vắng không phép"],
                counts["Chưa điểm danh"],
                attendance_rate,
            ]
        )
        student_rows.append((attendance_rate, student.HoTen.casefold(), row))

    for _, _, row in sorted(student_rows, key=lambda item: (item[0], item[1])):
        matrix_sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in matrix_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    matrix_sheet.freeze_panes = "D2"
    matrix_sheet.auto_filter.ref = matrix_sheet.dimensions
    matrix_sheet.column_dimensions["A"].width = 15
    matrix_sheet.column_dimensions["B"].width = 26
    matrix_sheet.column_dimensions["C"].width = 32
    for column_index in range(4, matrix_sheet.max_column + 1):
        matrix_sheet.column_dimensions[
            matrix_sheet.cell(row=1, column=column_index).column_letter
        ].width = 22 if column_index <= 3 + len(schedules) else 20
    matrix_sheet.row_dimensions[1].height = 42
    attendance_columns_end = 3 + len(schedules)
    for row_index in range(2, matrix_sheet.max_row + 1):
        for column_index in range(4, attendance_columns_end + 1):
            cell = matrix_sheet.cell(row=row_index, column=column_index)
            cell.fill = status_fills.get(cell.value, status_fills["—"])
            cell.alignment = Alignment(horizontal="center", vertical="center")
        matrix_sheet.cell(row=row_index, column=attendance_columns_end + 8).number_format = '0.0"%"'

    summary_sheet = workbook.create_sheet("Tổng hợp")
    summary_sheet.append(["Lớp", class_record.MaLop, class_record.TenLop])
    summary_sheet.append(["Ngày xuất", business_now().date().isoformat()])
    summary_sheet.append(["Số buổi đã lên lịch", len(schedules)])
    summary_sheet.append(["Tổng học viên", len(students)])
    summary_sheet.append(["Tổng lượt có mặt", totals["Có mặt"]])
    summary_sheet.append(["Tổng lượt đi muộn", totals["Đi muộn"]])
    summary_sheet.append(["Tổng lượt vắng có phép", totals["Vắng có phép"]])
    summary_sheet.append(["Tổng lượt vắng không phép", totals["Vắng không phép"]])
    summary_sheet.append(["Tổng lượt chưa điểm danh", totals["Chưa điểm danh"]])
    summary_sheet.append(
        [
            "Tổng lượt vắng",
            totals["Vắng có phép"] + totals["Vắng không phép"],
        ]
    )
    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 22
    summary_sheet.column_dimensions["C"].width = 30
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)

    output = BytesIO()
    workbook.save(output)
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", class_record.MaLop)
    filename = f"{safe_code}-chuyen-can-{business_now().date().isoformat()}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def decide_absence_request(
    request_id: int,
    decision: AbsenceDecision | AbsenceRequestDecision,
    current_user: NguoiDung,
    db: Session,
):
    if isinstance(decision, AbsenceRequestDecision):
        return decide_session_absence_request(
            request_id,
            decision,
            current_user,
            db,
        )
    return decide_class_absence_request(
        request_id,
        decision,
        current_user,
        db,
    )
