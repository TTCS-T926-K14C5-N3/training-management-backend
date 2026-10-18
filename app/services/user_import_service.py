from datetime import datetime, timedelta

import hashlib
import secrets
from io import BytesIO

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from pwdlib import PasswordHash
from app.core.config import FRONTEND_URL
from app.models.nguoi_dung import NguoiDung
from app.models.nguoi_dung_vai_tro import NguoiDungVaiTro
from app.models.token_kich_hoat import TokenKichHoat
from app.services.email_service import send_activation_email

password_hash = PasswordHash.recommended()
# ============================================================
# CẤU HÌNH
# ============================================================

ROLE_ID_TO_NAME = {
    1: "Guest",
    2: "Student",
    3: "Instructor",
    4: "TA",
    5: "Admissions",
    6: "Accountant",
    7: "Training Manager",
    8: "Admin",
}

ROLE_NAME_TO_ID = {
    "guest": 1,
    "student": 2,
    "instructor": 3,
    "ta": 4,
    "admissions": 5,
    "accountant": 6,
    "training manager": 7,
    "admin": 8,
}


# ============================================================
# HÀM HỖ TRỢ
# ============================================================

def normalize_text(value) -> str:
    """
    Chuyển dữ liệu Excel về chuỗi sạch:
    - None -> ""
    - bỏ khoảng trắng đầu/cuối
    - nhiều khoảng trắng liên tiếp -> 1 khoảng trắng
    """
    if value is None:
        return ""

    return " ".join(str(value).strip().split())


def normalize_email(value) -> str:
    """
    Chuẩn hóa email.
    """
    return normalize_text(value).lower()


def generate_temporary_password(length: int = 12) -> str:
    """
    Tạo mật khẩu tạm thời gồm:
    - chữ thường
    - chữ hoa
    - số
    """

    if length < 3:
        raise ValueError("Độ dài mật khẩu phải >= 3")

    lowercase = "abcdefghijklmnopqrstuvwxyz"
    uppercase = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    digits = "0123456789"

    password = [
        secrets.choice(lowercase),
        secrets.choice(uppercase),
        secrets.choice(digits),
    ]

    all_chars = lowercase + uppercase + digits

    for _ in range(length - 3):
        password.append(secrets.choice(all_chars))

    secrets.SystemRandom().shuffle(password)

    return "".join(password)


def generate_activation_token() -> tuple[str, str]:
    """
    Trả về:
    - raw token: dùng để gửi cho người dùng
    - token hash: lưu vào database
    """

    raw_token = secrets.token_urlsafe(32)

    token_hash = hashlib.sha256(
        raw_token.encode("utf-8")
    ).hexdigest()

    return raw_token, token_hash


def parse_role(value):
    """
    Excel cho phép nhập vai trò theo 2 dạng:

    1
    2
    ...
    8

    hoặc:

    Guest
    Student
    Instructor
    TA
    Admissions
    Accountant
    Training Manager
    Admin
    """

    if value is None:
        return None

    # Nếu Excel lưu số
    if isinstance(value, int):
        if 1 <= value <= 8:
            return value
        return None

    # Nếu Excel lưu số dạng float, ví dụ 3.0
    if isinstance(value, float):
        if value.is_integer():
            role_id = int(value)

            if 1 <= role_id <= 8:
                return role_id

        return None

    role_value = normalize_text(value)

    if not role_value:
        return None

    # Nhập "3"
    if role_value.isdigit():
        role_id = int(role_value)

        if 1 <= role_id <= 8:
            return role_id

        return None

    # Nhập tên vai trò
    return ROLE_NAME_TO_ID.get(role_value.lower())


def validate_excel_headers(headers):
    """
    Excel bắt buộc có đúng 3 cột:

    Họ tên
    Email
    Vai trò
    """

    normalized_headers = [
        normalize_text(header).lower()
        for header in headers
    ]

    required_headers = [
        "họ tên",
        "email",
        "vai trò",
    ]

    missing_headers = [
        header
        for header in required_headers
        if header not in normalized_headers
    ]

    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=(
                "File Excel thiếu cột bắt buộc: "
                + ", ".join(missing_headers)
                + ". Cần có: Họ tên, Email, Vai trò."
            ),
        )

    return {
        "ho_ten": normalized_headers.index("họ tên"),
        "email": normalized_headers.index("email"),
        "vai_tro": normalized_headers.index("vai trò"),
    }


def read_excel_file(file: UploadFile):
    """
    Đọc file .xlsx.

    Router hiện tại là def nên dùng:
        file.file.read()

    Không dùng await file.read().
    """

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Vui lòng chọn file Excel.",
        )

    filename = file.filename.lower()

    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Chỉ hỗ trợ file Excel .xlsx.",
        )

    try:
        content = file.file.read()

        if not content:
            raise HTTPException(
                status_code=400,
                detail="File Excel đang trống.",
            )

        workbook = load_workbook(
            filename=BytesIO(content),
            data_only=True,
        )

        worksheet = workbook.active

        return worksheet

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Không thể đọc file Excel: {str(exc)}",
        )


def get_role_id_from_database(
    db: Session,
    role_id: int,
):
    """
    Kiểm tra vai trò có tồn tại trong database hay không.
    Chỉ cho phép 8 vai trò hệ thống.
    """

    if role_id < 1 or role_id > 8:
        return None

    row = db.execute(
        text(
            """
            SELECT VaiTroId, TenVaiTro
            FROM VaiTro
            WHERE VaiTroId = :vai_tro_id
            """
        ),
        {
            "vai_tro_id": role_id,
        },
    ).fetchone()

    return row


def find_existing_email(
    db: Session,
    email: str,
):
    """
    Kiểm tra email đã tồn tại hay chưa.
    """

    return (
        db.query(NguoiDung)
        .filter(NguoiDung.Email == email)
        .first()
    )


# ============================================================
# ĐỌC VÀ VALIDATE CÁC DÒNG EXCEL
# ============================================================

def parse_excel_rows(
    db: Session,
    worksheet,
):
    """
    Đọc toàn bộ dòng Excel và trả về:

    {
        "valid_rows": [...],
        "invalid_rows": [...],
        "total_rows": ...
    }

    Không tạo tài khoản ở bước này.
    """

    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="File Excel không có dữ liệu.",
        )

    headers = rows[0]

    column_indexes = validate_excel_headers(
        headers
    )

    valid_rows = []
    invalid_rows = []

    # Dùng để phát hiện email trùng ngay trong Excel
    excel_emails = set()

    for excel_row_number, row in enumerate(
        rows[1:],
        start=2,
    ):
        ho_ten = normalize_text(
            row[column_indexes["ho_ten"]]
            if len(row) > column_indexes["ho_ten"]
            else ""
        )

        email = normalize_email(
            row[column_indexes["email"]]
            if len(row) > column_indexes["email"]
            else ""
        )

        vai_tro_raw = (
            row[column_indexes["vai_tro"]]
            if len(row) > column_indexes["vai_tro"]
            else ""
        )

        # ----------------------------------------------------
        # Bỏ qua dòng hoàn toàn trống
        # ----------------------------------------------------

        if not ho_ten and not email and not normalize_text(vai_tro_raw):
            continue

        errors = []

        # ----------------------------------------------------
        # Kiểm tra họ tên
        # ----------------------------------------------------

        if not ho_ten:
            errors.append("Thiếu họ tên.")

        elif len(ho_ten) > 100:
            errors.append(
                "Họ tên không được vượt quá 100 ký tự."
            )

        # ----------------------------------------------------
        # Kiểm tra email
        # ----------------------------------------------------

        if not email:
            errors.append("Thiếu email.")

        elif "@" not in email:
            errors.append("Email không hợp lệ.")

        elif len(email) > 255:
            errors.append(
                "Email không được vượt quá 255 ký tự."
            )

        # ----------------------------------------------------
        # Kiểm tra email trùng trong Excel
        # ----------------------------------------------------

        if email:
            if email in excel_emails:
                errors.append(
                    "Email bị trùng trong file Excel."
                )
            else:
                excel_emails.add(email)

        # ----------------------------------------------------
        # Kiểm tra email đã tồn tại trong DB
        # ----------------------------------------------------

        if email:
            existing_user = find_existing_email(
                db,
                email,
            )

            if existing_user:
                errors.append(
                    "Email đã tồn tại trong hệ thống."
                )

        # ----------------------------------------------------
        # Kiểm tra vai trò
        # ----------------------------------------------------

        role_id = parse_role(vai_tro_raw)

        if role_id is None:
            errors.append(
                "Vai trò không hợp lệ. "
                "Chỉ chấp nhận ID từ 1 đến 8 "
                "hoặc tên vai trò hợp lệ."
            )

        else:
            role = get_role_id_from_database(
                db,
                role_id,
            )

            if not role:
                errors.append(
                    f"Vai trò ID {role_id} không tồn tại trong database."
                )

        # ----------------------------------------------------
        # Kết quả dòng
        # ----------------------------------------------------

        row_result = {
            "row": excel_row_number,
            "hoTen": ho_ten,
            "email": email,
            "vaiTroId": role_id,
            "vaiTro": (
                ROLE_ID_TO_NAME.get(role_id)
                if role_id
                else None
            ),
        }

        if errors:
            row_result["trangThai"] = "Loi"
            row_result["errors"] = errors

            invalid_rows.append(row_result)

        else:
            row_result["trangThai"] = "HopLe"
            row_result["errors"] = []

            valid_rows.append(row_result)

    return {
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "total_rows": len(valid_rows) + len(invalid_rows),
    }


# ============================================================
# PREVIEW EXCEL
# ============================================================

def preview_user_import(
    db: Session,
    file: UploadFile,
):
    """
    Preview file Excel.

    Chỉ kiểm tra dữ liệu.
    KHÔNG tạo tài khoản.
    KHÔNG gửi email.
    """

    worksheet = read_excel_file(file)

    result = parse_excel_rows(
        db=db,
        worksheet=worksheet,
    )

    return {
        "message": "Đã kiểm tra file Excel.",
        "totalRows": result["total_rows"],
        "validRows": len(result["valid_rows"]),
        "invalidRows": len(result["invalid_rows"]),
        "canImport": len(result["valid_rows"]) > 0,
        "validData": result["valid_rows"],
        "invalidData": result["invalid_rows"],
    }


# ============================================================
# TẠO 1 TÀI KHOẢN TỪ DÒNG EXCEL
# ============================================================

def create_user_from_excel_row(
    db: Session,
    row: dict,
):
    """
    Tạo một tài khoản từ dòng Excel hợp lệ.

    Flow:

    1. Tạo mật khẩu tạm thời
    2. Hash mật khẩu
    3. Tạo NguoiDung
    4. Gán vai trò
    5. Tạo TokenKichHoat
    6. Gửi email
    """

    ho_ten = row["hoTen"]
    email = row["email"]
    role_id = row["vaiTroId"]

    # --------------------------------------------------------
    # Tạo mật khẩu tạm thời
    # --------------------------------------------------------

    temporary_password = generate_temporary_password(
        length=12
    )

    password_expiry = (
        datetime.now()
        + timedelta(hours=24)
    )

    # --------------------------------------------------------
    # Tạo user
    # --------------------------------------------------------

    now = datetime.now()

    user = NguoiDung(
        HoTen=ho_ten,
        Email=email,
        SoDienThoai=None,

        MatKhauHash=password_hash.hash(
            temporary_password
        ),

        TrangThai="HoatDong",

        DaKichHoat=False,

        BatBuocDoiMatKhau=True,

        CanDoiMatKhau=True,

        MatKhauTamThoiHetHan=password_expiry,

        SoLanDangNhapSai=0,

        KhoaDen=None,

        LyDoKhoa=None,

        NgayTao=now,
        NgayCapNhat=now,

        TokenValidAfter=None,
    )

    db.add(user)

    # Lấy NguoiDungId trước khi tạo quan hệ
    db.flush()

    # --------------------------------------------------------
    # Gán vai trò
    # --------------------------------------------------------

    user_role = NguoiDungVaiTro(
    NguoiDungId=user.NguoiDungId,
    VaiTroId=role_id,
    NgayGan=now,
)

    db.add(user_role)

    # --------------------------------------------------------
    # Tạo token kích hoạt
    # --------------------------------------------------------

    raw_token, token_hash = generate_activation_token()

    activation_expiry = (
        datetime.now()
        + timedelta(hours=24)
    )

    activation_token = TokenKichHoat(
        NguoiDungId=user.NguoiDungId,

        Token=token_hash,

        ThoiGianTao=now,

        ThoiGianHetHan=activation_expiry,

        DaSuDung=False,
    )

    db.add(activation_token)

    db.flush()

    # --------------------------------------------------------
    # Link kích hoạt / đăng nhập
    # --------------------------------------------------------

    activation_link = (
        f"{FRONTEND_URL.rstrip('/')}/activate-account?token={raw_token}"
    )

    # --------------------------------------------------------
    # Gửi email
    # --------------------------------------------------------

    send_activation_email(
        recipient_email=email,
        ho_ten=ho_ten,
        temporary_password=temporary_password,
        activation_link=activation_link,
    )

    return {
        "nguoiDungId": user.NguoiDungId,
        "hoTen": ho_ten,
        "email": email,
        "vaiTroId": role_id,
        "vaiTro": ROLE_ID_TO_NAME.get(role_id),
        "matKhauTamThoiHetHan": password_expiry,
    }


# ============================================================
# IMPORT CÁC DÒNG HỢP LỆ
# ============================================================

def import_valid_users(
    db: Session,
    file: UploadFile,
):
    """
    Import các dòng Excel hợp lệ.

    Dòng lỗi sẽ được bỏ qua.

    Ví dụ:

    10 dòng Excel
    - 7 dòng hợp lệ
    - 3 dòng lỗi

    -> tạo 7 tài khoản.
    """

    worksheet = read_excel_file(file)

    result = parse_excel_rows(
        db=db,
        worksheet=worksheet,
    )

    valid_rows = result["valid_rows"]
    invalid_rows = result["invalid_rows"]

    created_users = []
    failed_rows = []

    # --------------------------------------------------------
    # Không có dòng hợp lệ
    # --------------------------------------------------------

    if not valid_rows:
        return {
            "message": (
                "Không có dòng dữ liệu hợp lệ để import."
            ),
            "totalRows": result["total_rows"],
            "validRows": 0,
            "invalidRows": len(invalid_rows),
            "createdRows": 0,
            "failedRows": 0,
            "createdUsers": [],
            "invalidData": invalid_rows,
            "failedData": [],
        }

    # --------------------------------------------------------
    # Tạo từng tài khoản
    # --------------------------------------------------------

    for row in valid_rows:

        try:
            created_user = create_user_from_excel_row(
                db=db,
                row=row,
            )

            # Commit từng user.
            #
            # Nếu một dòng lỗi thì những dòng trước
            # vẫn giữ nguyên.
            db.commit()

            created_users.append(
                {
                    **created_user,
                    "row": row["row"],
                }
            )

        except Exception as exc:

            db.rollback()

            failed_rows.append(
                {
                    **row,
                    "trangThai": "Loi",
                    "errors": [
                        f"Không thể tạo tài khoản: {str(exc)}"
                    ],
                }
            )

    # --------------------------------------------------------
    # Kết quả
    # --------------------------------------------------------

    return {
        "message": (
            "Import Excel hoàn tất."
        ),

        "totalRows": result["total_rows"],

        "validRows": len(valid_rows),

        "invalidRows": len(invalid_rows),

        "createdRows": len(created_users),

        "failedRows": len(failed_rows),

        "createdUsers": created_users,

        "invalidData": invalid_rows,

        "failedData": failed_rows,
    }