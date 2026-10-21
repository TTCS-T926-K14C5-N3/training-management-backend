from io import BytesIO
from datetime import date, datetime, timedelta

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.database import Base
from app.models.attendance import (
    AbsenceRequest,
    AbsenceThreshold,
    AttendanceAudit,
    AttendanceNotification,
    AttendanceRecord,
)
from app.models.chuong_trinh_dao_tao import ChuongTrinhDaoTao
from app.models.diem_danh_buoi_hoc import DiemDanhBuoiHoc
from app.models.lich_hoc import LichHoc
from app.models.lop_hoc import LopHoc
from app.models.lop_hoc_hoc_vien import LopHocHocVien
from app.models.nguoi_dung import NguoiDung
from app.models.nguoi_dung_vai_tro import NguoiDungVaiTro
from app.routers.attendance import (
    AbsenceDecision,
    AbsenceRequestInput,
    AttendanceBatch,
    AttendanceItem,
    decide_absence_request,
    export_attendance,
    export_attendance_matrix,
    get_attendance_matrix,
    mark_all_present,
    save_attendance,
    submit_absence_request,
    ThresholdInput,
    update_threshold,
    router as attendance_router,
)
from app.main import app


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def user(db: Session, user_id: int, role_id: int, name: str) -> NguoiDung:
    person = NguoiDung(
        NguoiDungId=user_id,
        HoTen=name,
        Email=f"user{user_id}@example.test",
        MatKhauHash="test-hash",
        TrangThai="HoatDong",
        DaKichHoat=True,
        BatBuocDoiMatKhau=False,
        CanDoiMatKhau=False,
    )
    db.add(person)
    db.flush()
    db.add(NguoiDungVaiTro(NguoiDungId=user_id, VaiTroId=role_id, NgayGan=datetime.now()))
    db.flush()
    return person


def classroom(db: Session):
    teacher = user(db, 10, 3, "Teacher")
    student_a = user(db, 20, 2, "Student A")
    student_b = user(db, 21, 2, "Student B")
    manager = user(db, 30, 7, "Manager")
    program = ChuongTrinhDaoTao(
        ChuongTrinhId=40, MaChuongTrinh="PRG-40", TenChuongTrinh="Program",
        TongThoiLuong=20, HocPhiChuan=0, NgayTao=datetime.now(),
    )
    db.add(program)
    db.flush()
    cls = LopHoc(
        LopHocId=50, MaLop="ATT-50", TenLop="Attendance class", NamHoc="2026-2027",
        ChuongTrinhId=40, GiangVienId=teacher.NguoiDungId, NgayTao=datetime.now(),
    )
    db.add(cls)
    db.flush()
    for person in (student_a, student_b):
        db.add(LopHocHocVien(LopHocId=cls.LopHocId, HocVienId=person.NguoiDungId, NgayThamGia=datetime.now()))
    db.add(AbsenceThreshold(ChuongTrinhId=40, TiLeToiDa=20, NgayCapNhat=datetime.now()))
    db.commit()
    return teacher, student_a, student_b, manager, cls


def test_mark_all_present_creates_auditable_records(db):
    teacher, _, _, _, cls = classroom(db)
    response = mark_all_present(cls.LopHocId, date.today(), teacher, db)
    assert response["counts"] == {
        "present": 2,
        "late": 0,
        "excused": 0,
        "unexcused": 0,
        "absent": 0,
        "unmarked": 0,
        "total": 2,
    }
    assert db.query(AttendanceRecord).count() == 2
    assert db.query(AttendanceAudit).count() == 2


def test_attendance_export_is_a_real_xlsx_with_summary_and_details(db):
    teacher, _, _, _, cls = classroom(db)
    response = export_attendance(cls.LopHocId, date.today(), teacher, db)
    workbook = load_workbook(BytesIO(response.body), read_only=True)
    assert workbook.sheetnames == ["Tổng hợp", "Chi tiết"]
    assert workbook["Tổng hợp"]["A3"].value == "Tổng học viên"
    assert workbook["Chi tiết"]["A1"].value == "Mã học viên"
    assert response.headers["content-disposition"].endswith(
        f'filename="{cls.MaLop}-{date.today().isoformat()}.xlsx"'
    )


def test_attendance_matrix_export_contains_session_columns_and_student_totals(
    db,
    monkeypatch,
):
    fixed_now = datetime.combine(date.today(), datetime.max.time())
    monkeypatch.setattr(
        "app.routers.attendance.business_now_naive",
        lambda: fixed_now,
    )
    monkeypatch.setattr(
        "app.routers.attendance.business_now",
        lambda: fixed_now,
    )
    teacher, student_a, student_b, _, cls = classroom(db)
    previous_day = date.today() - timedelta(days=1)
    for enrollment in db.query(LopHocHocVien).filter(
        LopHocHocVien.LopHocId == cls.LopHocId
    ):
        enrollment.NgayThamGia = datetime.combine(
            previous_day - timedelta(days=10),
            datetime.min.time(),
        )

    first_session = LichHoc(
        LopHocId=cls.LopHocId,
        MonHocId=1,
        NgayHoc=previous_day,
        GioBatDau=datetime.strptime("08:00", "%H:%M").time(),
        GioKetThuc=datetime.strptime("10:00", "%H:%M").time(),
        PhongHocId=1,
        GiangVienId=teacher.NguoiDungId,
        TrangThai="Đã lên lịch",
        NgayTao=datetime.now(),
        NguoiTaoId=teacher.NguoiDungId,
        NgayCapNhat=datetime.now(),
        NguoiCapNhatId=teacher.NguoiDungId,
    )
    second_session = LichHoc(
        LopHocId=cls.LopHocId,
        MonHocId=1,
        NgayHoc=date.today(),
        GioBatDau=datetime.strptime("10:00", "%H:%M").time(),
        GioKetThuc=datetime.strptime("12:00", "%H:%M").time(),
        PhongHocId=1,
        GiangVienId=teacher.NguoiDungId,
        TrangThai="Đã lên lịch",
        NgayTao=datetime.now(),
        NguoiTaoId=teacher.NguoiDungId,
        NgayCapNhat=datetime.now(),
        NguoiCapNhatId=teacher.NguoiDungId,
    )
    future_session = LichHoc(
        LopHocId=cls.LopHocId,
        MonHocId=1,
        NgayHoc=date.today() + timedelta(days=1),
        GioBatDau=datetime.strptime("08:00", "%H:%M").time(),
        GioKetThuc=datetime.strptime("10:00", "%H:%M").time(),
        PhongHocId=1,
        GiangVienId=teacher.NguoiDungId,
        TrangThai="Đã lên lịch",
        NgayTao=datetime.now(),
        NguoiTaoId=teacher.NguoiDungId,
        NgayCapNhat=datetime.now(),
        NguoiCapNhatId=teacher.NguoiDungId,
    )
    db.add_all([first_session, second_session, future_session])
    db.flush()
    db.add_all(
        [
            DiemDanhBuoiHoc(
                LichHocId=first_session.LichHocId,
                HocVienId=student_a.NguoiDungId,
                TrangThai="Có mặt",
                NgayCapNhat=datetime.now(),
                NguoiCapNhatId=teacher.NguoiDungId,
            ),
            DiemDanhBuoiHoc(
                LichHocId=second_session.LichHocId,
                HocVienId=student_a.NguoiDungId,
                TrangThai="Đi muộn",
                NgayCapNhat=datetime.now(),
                NguoiCapNhatId=teacher.NguoiDungId,
            ),
            DiemDanhBuoiHoc(
                LichHocId=first_session.LichHocId,
                HocVienId=student_b.NguoiDungId,
                TrangThai="Vắng không phép",
                NgayCapNhat=datetime.now(),
                NguoiCapNhatId=teacher.NguoiDungId,
            ),
        ]
    )
    db.commit()

    response = export_attendance_matrix(cls.LopHocId, teacher, db)
    workbook = load_workbook(BytesIO(response.body))
    matrix = workbook["Ma trận chuyên cần"]

    assert workbook.sheetnames == ["Ma trận chuyên cần", "Tổng hợp"]
    assert matrix["D1"].value.startswith(f"BH-{first_session.LichHocId:06d}")
    assert matrix["E1"].value.startswith(f"BH-{second_session.LichHocId:06d}")
    assert "BH-" not in matrix["F1"].value
    assert matrix["M1"].value == "Tỷ lệ chuyên cần (%)"
    assert matrix["D2"].value == "Vắng không phép"
    assert matrix["F2"].value == 2
    assert matrix["G2"].value == 1
    assert matrix["D2"].fill.fgColor.rgb.endswith("FEE2E2")
    assert matrix["D3"].value == "Có mặt"
    assert matrix["E3"].value == "Đi muộn"
    assert matrix["F3"].value == 2
    assert matrix["G3"].value == 0
    assert matrix["H3"].value == 1
    assert matrix["I3"].value == 1
    assert matrix["M3"].value == 100
    assert matrix["L2"].value == 1
    assert matrix["B2"].value == "Student B"
    assert matrix["B3"].value == "Student A"
    assert matrix["E2"].value == "Chưa điểm danh"
    assert matrix["L1"].value == "Chưa điểm danh"
    assert matrix["N1"].value is None
    assert workbook["Tổng hợp"]["B3"].value == 2
    assert response.headers["content-disposition"].endswith(
        f'filename="{cls.MaLop}-chuyen-can-{date.today().isoformat()}.xlsx"'
    )


def test_threshold_crossing_flags_and_notifies_teacher_and_manager(db):
    teacher, student, _, manager, cls = classroom(db)
    save_attendance(cls.LopHocId, AttendanceBatch(date=date.today(), records=[
        AttendanceItem(studentId=student.NguoiDungId, status="unexcused_absence"),
    ]), teacher, db)
    notices = db.query(AttendanceNotification).filter(
        AttendanceNotification.HocVienId == student.NguoiDungId,
        AttendanceNotification.Loai == "absence_threshold",
    ).all()
    assert {notice.NguoiNhanId for notice in notices} == {
        teacher.NguoiDungId,
        manager.NguoiDungId,
        student.NguoiDungId,
    }
    matrix = get_attendance_matrix(cls.LopHocId, date.today(), teacher, db)
    assert matrix["aboveThreshold"] is True


def test_only_manager_can_edit_past_attendance(db):
    teacher, student, _, manager, cls = classroom(db)
    past_day = date.today() - timedelta(days=1)
    enrollment = db.query(LopHocHocVien).filter(
        LopHocHocVien.LopHocId == cls.LopHocId,
        LopHocHocVien.HocVienId == student.NguoiDungId,
    ).one()
    enrollment.NgayThamGia = datetime.combine(past_day - timedelta(days=1), datetime.min.time())
    db.commit()
    payload = AttendanceBatch(date=past_day, records=[
        AttendanceItem(studentId=student.NguoiDungId, status="present"),
    ])
    with pytest.raises(HTTPException) as error:
        save_attendance(cls.LopHocId, payload, teacher, db)
    assert error.value.status_code == 403
    response = save_attendance(cls.LopHocId, payload, manager, db)
    assert response["date"] == payload.attendance_date.isoformat()


def test_approved_absence_request_marks_excused_and_notifies_student(db):
    teacher, student, _, _, cls = classroom(db)
    created = submit_absence_request(
        AbsenceRequestInput(classId=cls.LopHocId, date=date.today(), reason="Medical appointment"),
        student,
        db,
    )
    result = decide_absence_request(
        created["requestId"], AbsenceDecision(decision="approved", reason="Approved"), teacher, db,
    )
    record = db.query(AttendanceRecord).filter(AttendanceRecord.HocVienId == student.NguoiDungId).one()
    assert result["status"] == "approved"
    assert record.TrangThai == "excused_absence"
    assert db.query(AttendanceNotification).filter(
        AttendanceNotification.NguoiNhanId == student.NguoiDungId,
        AttendanceNotification.Loai == "absence_request_approved",
    ).count() == 1
    assert db.query(AttendanceAudit).filter(AttendanceAudit.DoiTuong == "absence_request").count() == 2


def test_lowering_program_threshold_notifies_existing_flagged_student(db):
    teacher, student, _, manager, cls = classroom(db)
    threshold = db.query(AbsenceThreshold).filter(AbsenceThreshold.ChuongTrinhId == cls.ChuongTrinhId).one()
    threshold.TiLeToiDa = 60
    db.add_all([
        AttendanceRecord(LopHocId=cls.LopHocId, HocVienId=student.NguoiDungId,
                         NgayHoc=date.today() - timedelta(days=2), TrangThai="unexcused_absence",
                         NguoiCapNhatId=teacher.NguoiDungId, NgayCapNhat=datetime.now()),
        AttendanceRecord(LopHocId=cls.LopHocId, HocVienId=student.NguoiDungId,
                         NgayHoc=date.today() - timedelta(days=1), TrangThai="present",
                         NguoiCapNhatId=teacher.NguoiDungId, NgayCapNhat=datetime.now()),
    ])
    db.commit()
    update_threshold(cls.ChuongTrinhId, ThresholdInput(percentage=40), manager, db)
    assert {row.NguoiNhanId for row in db.query(AttendanceNotification).filter(
        AttendanceNotification.HocVienId == student.NguoiDungId,
        AttendanceNotification.Loai == "absence_threshold",
    ).all(    )} == {teacher.NguoiDungId, manager.NguoiDungId, student.NguoiDungId}




def test_attendance_routes_cover_workflows():
    paths = {route.path for route in attendance_router.routes}
    assert "/attendance/classes/{class_id}/matrix" in paths
    assert "/attendance/classes/{class_id}/mark-all-present" in paths
    assert "/attendance/absence-requests/{request_id}/decision" in paths
    assert "/attendance/my-history" in paths
    assert "/attendance/classes/{class_id}/export" in paths
    assert "/attendance/classes/{class_id}/export-matrix" in paths
    assert "/attendance/notifications" in paths

    app_paths = set(app.openapi()["paths"])
    assert "/attendance/absence-requests" in app_paths
    assert "/attendance/sessions/absence-requests" in app_paths
    assert "/attendance/sessions/absence-requests/with-evidence" in app_paths
